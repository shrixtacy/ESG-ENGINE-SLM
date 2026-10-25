"""
ESGPilot — Table Parser.

Handles tabular ESG data extraction and normalization.
Supports:
  - BRSR-style disclosure tables (rows = indicators, cols = FY values)
  - Emission factor tables (CSV/Excel)
  - Peer comparison tables
  - Both row-level chunks and full-table Markdown for embedding
"""

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Optional, Any


@dataclass
class TableData:
    """Parsed tabular data."""
    headers: list[str]
    rows: list[list[Any]]
    source: str = ""
    section_path: Optional[str] = None
    framework: Optional[str] = None
    note: Optional[str] = None

    def to_markdown(self) -> str:
        """Convert to Markdown table string."""
        if not self.headers:
            return ""
        sep = ["---"] * len(self.headers)
        lines = [
            "| " + " | ".join(str(h) for h in self.headers) + " |",
            "| " + " | ".join(sep) + " |",
        ]
        for row in self.rows:
            # Pad row if shorter than headers
            padded = list(row) + [""] * (len(self.headers) - len(row))
            lines.append("| " + " | ".join(str(c) for c in padded[:len(self.headers)]) + " |")
        return "\n".join(lines)

    def to_row_sentences(self) -> list[str]:
        """
        Convert each row into a natural-language sentence for embedding.
        Useful for fine-grained retrieval of individual data points.
        """
        sentences = []
        for row in self.rows:
            parts = []
            for header, value in zip(self.headers, row):
                if value not in ("", None, "N/A", "-"):
                    parts.append(f"{header}: {value}")
            if parts:
                sentence = ". ".join(parts) + "."
                sentences.append(sentence)
        return sentences

    def to_chunks(self, rows_per_chunk: int = 10) -> list[str]:
        """
        Split table into chunks of `rows_per_chunk` rows.
        Each chunk includes headers for context.
        """
        chunks = []
        header_md = "| " + " | ".join(str(h) for h in self.headers) + " |"
        sep_md = "| " + " | ".join(["---"] * len(self.headers)) + " |"

        for i in range(0, len(self.rows), rows_per_chunk):
            batch = self.rows[i:i + rows_per_chunk]
            lines = [header_md, sep_md]
            for row in batch:
                padded = list(row) + [""] * (len(self.headers) - len(row))
                lines.append("| " + " | ".join(str(c) for c in padded[:len(self.headers)]) + " |")
            chunks.append("\n".join(lines))

        return chunks


class TableParser:
    """
    Multi-format table parser.
    Handles CSV, Excel (via openpyxl), and raw list-of-lists input.
    """

    # ──────────────────────────────────────────────
    # CSV
    # ──────────────────────────────────────────────

    @staticmethod
    def parse_csv(
        filepath_or_content: str,
        is_content: bool = False,
        encoding: str = "utf-8",
        framework: Optional[str] = None,
        section_path: Optional[str] = None,
    ) -> TableData:
        """
        Parse a CSV file or CSV string content.

        Args:
            filepath_or_content: Path to CSV file OR raw CSV string (if is_content=True)
            is_content: If True, treat first arg as CSV string
        """
        if is_content:
            reader = csv.reader(io.StringIO(filepath_or_content))
        else:
            f = open(filepath_or_content, "r", encoding=encoding, newline="")
            reader = csv.reader(f)

        rows = list(reader)

        if is_content:
            pass
        else:
            f.close()

        if not rows:
            return TableData(headers=[], rows=[])

        headers = [h.strip() for h in rows[0]]
        data_rows = [
            [cell.strip() for cell in row]
            for row in rows[1:]
            if any(cell.strip() for cell in row)
        ]

        return TableData(
            headers=headers,
            rows=data_rows,
            source=filepath_or_content if not is_content else "inline",
            framework=framework,
            section_path=section_path,
        )

    @staticmethod
    def parse_excel(
        filepath: str,
        sheet_name: Optional[str] = None,
        header_row: int = 0,
        framework: Optional[str] = None,
        section_path: Optional[str] = None,
    ) -> list[TableData]:
        """
        Parse an Excel file, returning one TableData per sheet (or target sheet).
        Requires openpyxl.
        """
        try:
            import openpyxl
        except ImportError:
            return [TableData(
                headers=[], rows=[],
                note="openpyxl not installed. Run: pip install openpyxl",
            )]

        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        tables = []

        sheet_names = [sheet_name] if sheet_name else wb.sheetnames

        for name in sheet_names:
            if name not in wb.sheetnames:
                continue

            ws = wb[name]
            all_rows = list(ws.iter_rows(values_only=True))

            if not all_rows or header_row >= len(all_rows):
                continue

            headers = [str(c).strip() if c is not None else "" for c in all_rows[header_row]]

            data_rows = []
            for row in all_rows[header_row + 1:]:
                row_data = [c if c is not None else "" for c in row]
                # Skip fully empty rows
                if any(str(c).strip() for c in row_data):
                    data_rows.append([str(c) for c in row_data])

            tables.append(TableData(
                headers=headers,
                rows=data_rows,
                source=f"{filepath}::{name}",
                framework=framework,
                section_path=section_path or name,
            ))

        wb.close()
        return tables

    # ──────────────────────────────────────────────
    # Emission Factor Table
    # ──────────────────────────────────────────────

    @staticmethod
    def parse_emission_factor_table(rows: list[list]) -> list[dict]:
        """
        Parse a raw emission factor table (list of lists) into structured records.
        Expected columns: source, year, region, fuel_or_grid, unit, factor, gwp_basis
        """
        records = []
        for row in rows:
            if len(row) < 6:
                continue
            try:
                record = {
                    "source": str(row[0]).strip(),
                    "year": int(row[1]),
                    "region": str(row[2]).strip(),
                    "fuel_or_grid": str(row[3]).strip(),
                    "unit": str(row[4]).strip(),
                    "factor": float(str(row[5]).replace(",", "")),
                    "gwp_basis": str(row[6]).strip() if len(row) > 6 else "AR5",
                }
                records.append(record)
            except (ValueError, IndexError):
                continue
        return records

    # ──────────────────────────────────────────────
    # BRSR Table Normalizer
    # ──────────────────────────────────────────────

    @staticmethod
    def normalize_brsr_table(table: TableData) -> list[dict]:
        """
        Normalize a BRSR-style disclosure table.
        Typical shape: rows are parameters/indicators, cols are FY Current / FY Previous.

        Returns a list of normalized records with indicator_id, label, fy, value, unit.
        """
        records = []
        headers = table.headers

        # Detect FY columns (look for FY/Year patterns)
        fy_pattern = re.compile(r"(FY\s*\d{2,4}|\d{4}[-–]\d{2,4})", re.IGNORECASE)
        fy_cols = {
            i: headers[i]
            for i in range(len(headers))
            if fy_pattern.search(headers[i])
        }

        # Detect parameter/label column (usually first col)
        label_col = 0

        for row in table.rows:
            if not row or not row[label_col]:
                continue

            label = str(row[label_col]).strip()
            if not label:
                continue

            # Try to extract BRSR indicator ID from label
            indicator_id = None
            for pattern in [
                re.compile(r"(P\d[-\s]?[EL]I[-\s]?\d+[a-z]?)", re.IGNORECASE),
                re.compile(r"(GRI\s+\d{3}-\d+)", re.IGNORECASE),
            ]:
                m = pattern.search(label)
                if m:
                    indicator_id = m.group(1)
                    break

            for col_idx, fy_label in fy_cols.items():
                if col_idx >= len(row):
                    continue
                value = str(row[col_idx]).strip()
                if value in ("", "-", "N/A", "NA"):
                    continue

                # Try to detect unit in label
                unit = None
                unit_match = re.search(r"\(([^)]+)\)", label)
                if unit_match:
                    unit = unit_match.group(1)

                records.append({
                    "indicator_id": indicator_id,
                    "label": label,
                    "fy": fy_label,
                    "value": value,
                    "unit": unit,
                    "section_path": table.section_path,
                    "framework": table.framework,
                })

        return records
